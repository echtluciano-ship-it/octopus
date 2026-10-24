from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from service import delivery, export_excel, gmail


class DeliveryTests(unittest.TestCase):
    def setUp(self):
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        self.root = Path(self.stack.enter_context(tempfile.TemporaryDirectory()))
        self.rows = []
        for i, group in enumerate(("a", "a", "b"), 1):
            self.rows.append(dict(document_id=i, status="READY_FOR_SHEET",
                attachment_index=0, original_filename="test.png", source_system="test",
                source_document_id=None, source_url=None, source_batch="batch",
                arrival_sequence=i, extraction_id=i, group_id=group, group_name=group,
                document_type="transfer", ticket_date=None, amount_cents=i*100,
                amount=i, identifier_type="operation", identifier_value=str(i),
                full_name="Test", tax_id=None, origin_bank="Test", destination_account_key="test",
                observations="", review_reason="READY_FOR_SHEET"))
        self.stack.enter_context(patch.object(export_excel, "settings", SimpleNamespace(
            export_dir=self.root, group_name=lambda group, fallback=None: group)))
        self.loader = self.stack.enter_context(patch.object(export_excel, "load_delivery_rows", return_value=self.rows))
        self.integrity = self.stack.enter_context(patch.object(delivery, "receipt_integrity", return_value={"ready": True}))
        self.report = export_excel.generate(source_batch="batch")
        self.paths = [Path(b["path"]) for b in self.report["erika_workbooks"]]
        self.manifest = self.paths[0].parent / "manifest.json"
        self.conn = MagicMock()
        self.conn.execute.return_value.fetchall.return_value = [dict(document_id=i,
            object_key=str(i), sha256=hashlib.sha256(b"original").hexdigest(), size_bytes=8)
            for i in range(1, 4)]
        self.stack.enter_context(patch.object(delivery, "connection")).return_value.__enter__.return_value = self.conn
        self.original = self.stack.enter_context(patch.object(delivery, "get_original", return_value=b"original"))
        self.stack.enter_context(patch.object(gmail, "settings", SimpleNamespace(
            gmail_attachment_root=self.root, gmail_allowed_recipients=("test@example.com",))))

    def validate(self):
        return delivery.validate_delivery(self.manifest, self.paths)

    def write_report(self):
        self.manifest.write_text(json.dumps(self.report), encoding="utf-8")

    def test_two_groups_get_exactly_two_editable_workbooks(self):
        self.assertEqual(len(self.paths), 2)
        self.assertTrue(self.validate()["delivery_ready"])

    def test_missing_ingress_blocks_even_with_all_stored_rows_exported(self):
        self.integrity.return_value = {"ready": False}
        with self.assertRaises(ValueError):
            self.validate()
        report = export_excel.generate(source_batch="batch")
        self.assertFalse(report["delivery_gate"]["delivery_ready"])

    def test_nonterminal_state_after_export_blocks(self):
        self.rows[0]["status"] = "PROCESSING"
        with self.assertRaises(ValueError):
            self.validate()

    def test_missing_extraction_blocks(self):
        self.rows[0]["extraction_id"] = None
        with self.assertRaises(ValueError):
            self.validate()

    def test_missing_workbook_blocks(self):
        self.paths.pop()
        with self.assertRaises(ValueError):
            self.validate()

    def test_extra_internal_workbook_blocks(self):
        self.paths.append(self.manifest.parent / "CONTROL_LOTE.xlsx")
        with self.assertRaises(ValueError):
            self.validate()

    def test_reordered_excel_blocks(self):
        book = load_workbook(self.paths[0])
        sheet = book.active
        first, second = list(sheet.values)[1:]
        for col, (a, b) in enumerate(zip(first, second), 1):
            sheet.cell(2, col, b)
            sheet.cell(3, col, a)
        book.save(self.paths[0])
        book.close()
        with self.assertRaises(ValueError):
            self.validate()

    def test_deleted_row_blocks(self):
        book = load_workbook(self.paths[0])
        book.active.delete_rows(2)
        book.save(self.paths[0])
        book.close()
        with self.assertRaises(ValueError):
            self.validate()

    def test_artificial_duplicate_row_blocks(self):
        book = load_workbook(self.paths[0])
        book.active.append(list(book.active.values)[1])
        book.save(self.paths[0])
        book.close()
        with self.assertRaises(ValueError):
            self.validate()

    def test_protected_excel_blocks(self):
        book = load_workbook(self.paths[0])
        book.active.protection.sheet = True
        book.save(self.paths[0])
        book.close()
        with self.assertRaises(ValueError):
            self.validate()

    def test_changed_original_blocks(self):
        self.original.return_value = b"tampered"
        with self.assertRaises(ValueError):
            self.validate()

    def test_missing_original_blocks(self):
        self.original.side_effect = FileNotFoundError()
        with self.assertRaises(FileNotFoundError):
            self.validate()

    def test_forged_manifest_trace_blocks(self):
        self.report["erika_row_trace"][0]["document_id"] = 999
        self.write_report()
        with self.assertRaises(ValueError):
            self.validate()

    def test_queue_blocks_before_any_outbox_write(self):
        self.integrity.return_value = {"ready": False}
        with patch.object(gmail, "connection") as conn:
            with self.assertRaises(ValueError):
                gmail.queue_email(recipient="test@example.com", subject="test", body="test",
                                  attachments=self.paths, idempotency_key="test")
            conn.assert_not_called()

    def test_queue_requires_manifest(self):
        self.manifest.unlink()
        with patch.object(gmail, "connection") as conn:
            with self.assertRaises(FileNotFoundError):
                gmail.queue_email(recipient="test@example.com", subject="test", body="test",
                                  attachments=self.paths, idempotency_key="test")
            conn.assert_not_called()

    def test_send_revalidates_and_never_contacts_provider_when_blocked(self):
        payload = dict(recipient="test@example.com", subject="test", body="test",
            attachments=[gmail.attachment_evidence(p) for p in self.paths],
            delivery_manifest=gmail.attachment_evidence(self.manifest))
        conn = MagicMock()
        conn.execute.return_value.fetchone.return_value = dict(payload=payload,
            idempotency_key="test", delivery_status="NOT_SENT")
        self.integrity.return_value = {"ready": False}
        provider = MagicMock()
        with patch.object(gmail, "connection") as factory:
            factory.return_value.__enter__.return_value = conn
            with self.assertRaises(ValueError):
                gmail.send_action(1, service=provider)
        provider.users.assert_not_called()
        conn.commit.assert_not_called()

    def test_valid_queue_then_send_uses_only_mock_provider(self):
        conn = MagicMock()
        conn.execute.return_value.fetchone.return_value = {"action_id": 7}
        with patch.object(gmail, "connection") as factory:
            factory.return_value.__enter__.return_value = conn
            action = gmail.queue_email(recipient="test@example.com", subject="test", body="test",
                attachments=self.paths, idempotency_key="test")
            self.assertEqual(action, 7)
            payload = json.loads(conn.execute.call_args.args[1][1])
            conn.execute.return_value.fetchone.return_value = dict(payload=payload,
                idempotency_key="test", delivery_status="NOT_SENT")
            provider = MagicMock()
            provider.users.return_value.messages.return_value.send.return_value.execute.return_value = {"id": "mock"}
            self.assertEqual(gmail.send_action(7, service=provider)["status"], "SENT")
            provider.users.return_value.messages.return_value.send.assert_called_once()

    def test_legacy_email_without_manifest_is_blocked(self):
        conn = MagicMock()
        conn.execute.return_value.fetchone.return_value = dict(
            payload={"recipient": "test@example.com"}, delivery_status="NOT_SENT")
        provider = MagicMock()
        with patch.object(gmail, "connection") as factory:
            factory.return_value.__enter__.return_value = conn
            with self.assertRaises(ValueError):
                gmail.send_action(1, service=provider)
        provider.users.assert_not_called()
        conn.commit.assert_not_called()


class ReceiptIntegrityTests(unittest.TestCase):
    def check(self, rows, failures=0):
        with patch.object(delivery, "connection") as factory:
            conn = factory.return_value.__enter__.return_value
            conn.execute.return_value.fetchall.return_value = rows
            conn.execute.return_value.fetchone.return_value = {"count": failures}
            return delivery.receipt_integrity()

    def test_message_without_any_original_is_not_invisible(self):
        self.assertFalse(self.check([dict(message_id=1, expected_attachments=1,
            stored=0, invalid_indices=0)])["ready"])

    def test_equal_totals_cannot_hide_missing_and_extra_originals(self):
        self.assertFalse(self.check([
            dict(message_id=1, expected_attachments=1, stored=0, invalid_indices=0),
            dict(message_id=2, expected_attachments=1, stored=2, invalid_indices=1),
        ])["ready"])

    def test_unresolved_ingest_error_blocks(self):
        self.assertFalse(self.check([], failures=1)["ready"])

    def test_complete_receipts_pass(self):
        self.assertTrue(self.check([dict(message_id=1, expected_attachments=1,
            stored=1, invalid_indices=0)])["ready"])
