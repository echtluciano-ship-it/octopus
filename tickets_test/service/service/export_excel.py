from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from service.config import settings
from service.db import connection


TRANSFER_COLUMNS = ("FECHA TICHET", "MONTO", "NRO COMPROBANTE", "NOMBRE Y APELLIDO", "CUIT", "BANCO")
DEPOSIT_COLUMNS = ("FECHA", "IMPORTE", "NUMERO DE COMPROBANTE", "BANCO")
CONTROL_COLUMNS = (
    "DOCUMENT ID", "ORDEN", "ESTADO", "TIPO", "FECHA", "MONTO", "TIPO IDENTIFICADOR",
    "IDENTIFICADOR", "NOMBRE Y APELLIDO", "CUIT", "BANCO", "CUENTA DESTINO",
    "RESULTADO / MOTIVO", "ARCHIVO ORIGINAL", "DRIVE ID", "LINK DRIVE",
    "EXCEL", "FILA",
)


def _safe(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_") or "sin_cuenta"


def _prepare_output(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)
    for pattern in ("*.xlsx", "manifest.json"):
        for existing in path.glob(pattern):
            existing.unlink()


def _deposit_group_key(origin_bank: str | None) -> str:
    plain = re.sub(r"[^a-z0-9]", "", (origin_bank or "").lower())
    if "macro" in plain:
        return "macro"
    if "comafi" in plain:
        return "comafi"
    return "sin_banco_confirmado"


def _write(path: Path, rows: list[dict], *, deposit: bool = False, review: bool = False) -> None:
    columns = DEPOSIT_COLUMNS if deposit else TRANSFER_COLUMNS
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hoja1"
    sheet.append(columns)
    for cell in sheet[1]:
        cell.alignment = Alignment(horizontal="center")
    for item in rows:
        if deposit:
            values = (item["ticket_date"], item["amount"], item["identifier_value"], item["origin_bank"])
        else:
            values = (
                item["ticket_date"], item["amount"], item["identifier_value"],
                item["full_name"], item["tax_id"], item["origin_bank"],
            )
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(horizontal="center")
        if review:
            for cell in sheet[sheet.max_row]:
                cell.fill = PatternFill("solid", fgColor="FFD9D9")
                cell.comment = Comment(item["review_reason"], "OCTOPUS TEST")
    widths = (12, 18.71, 34, 30.71, 24.86, 22.29)
    for index, width in enumerate(widths[: len(columns)], start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for cell in sheet["A"][1:]:
        cell.number_format = "mm-dd-yy"
    for cell in sheet["B"][1:]:
        cell.number_format = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _write_control(path: Path, rows: list[dict], trace_by_document: dict[int, dict]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Control"
    sheet.append(CONTROL_COLUMNS)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    status_fill = {
        "READY_FOR_SHEET": "E2F0D9",
        "REVIEW": "FFF2CC",
        "WRONG_ACCOUNT_REVIEW": "FCE4D6",
        "CHECK_REVIEW": "FFF2CC",
        "DUPLICATE_CONFIRMED": "D9E1F2",
        "ERROR_FINAL": "F4CCCC",
        "ERROR_RETRYABLE": "F4CCCC",
    }
    for item in rows:
        trace = trace_by_document.get(item["document_id"], {})
        sheet.append((
            item["document_id"], item["arrival_sequence"], item["status"], item["document_type"],
            item["ticket_date"], item["amount"], item["identifier_type"],
            item["identifier_value"], item["full_name"], item["tax_id"],
            item["origin_bank"], item["destination_account_key"],
            item["review_reason"], item["original_filename"],
            item["source_document_id"], item["source_url"],
            trace.get("workbook"), trace.get("row"),
        ))
        row_number = sheet.max_row
        fill = status_fill.get(item["status"])
        if fill:
            sheet.cell(row_number, 3).fill = PatternFill("solid", fgColor=fill)
        if item["source_url"]:
            link_cell = sheet.cell(row_number, 16)
            link_cell.hyperlink = item["source_url"]
            link_cell.style = "Hyperlink"
    widths = (12, 9, 24, 14, 12, 16, 20, 38, 30, 18, 20, 22, 54, 44, 34, 18, 38, 8)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index) if index <= 26 else "A"].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:R{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    for cell in sheet["E"][1:]:
        cell.number_format = "dd/mm/yyyy"
    for cell in sheet["F"][1:]:
        cell.number_format = '$#,##0.00'
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def generate(
    source_batch: str | None = None,
    document_id: int | list[int] | None = None,
) -> dict:
    document_ids = (
        [document_id]
        if isinstance(document_id, int)
        else list(document_id or [])
    )
    if source_batch and document_ids:
        raise ValueError("source_batch and document_id are mutually exclusive")
    if source_batch:
        where_filter = "AND d.source_batch=%s"
        params = (source_batch,)
    elif document_ids:
        where_filter = "AND d.document_id = ANY(%s)"
        params = (document_ids,)
    else:
        where_filter = ""
        params = ()
    with connection() as conn:
        rows = conn.execute(
            f"""
            SELECT d.document_id,d.status,d.original_filename,d.source_system,
                   d.source_document_id,d.source_url,d.source_batch,m.arrival_sequence,e.document_type,e.ticket_date,
                   e.amount_cents,e.identifier_type,e.identifier_value,e.full_name,e.tax_id,e.origin_bank,
                   e.destination_account_key,
                   CASE WHEN d.status='DUPLICATE_CONFIRMED'
                        THEN COALESCE(d.duplicate_reason,'Duplicado confirmado.')
                        ELSE e.observations END AS observations
            FROM ticket_documents d
            JOIN inbound_messages m USING(message_id)
            JOIN LATERAL (
                SELECT * FROM ticket_extractions x
                WHERE x.document_id=COALESCE(d.duplicate_of,d.document_id)
                ORDER BY attempt DESC LIMIT 1
            ) e ON true
            WHERE d.status IN (
                'READY_FOR_SHEET','REVIEW','WRONG_ACCOUNT_REVIEW','CHECK_REVIEW',
                'DUPLICATE_CONFIRMED'
            )
              {where_filter}
            ORDER BY m.arrival_sequence,d.attachment_index
            """,
            params,
        ).fetchall()
    groups: dict[str, list[dict]] = defaultdict(list)
    review_rows: list[dict] = []
    deposits: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        item = dict(row)
        item["amount"] = item["amount_cents"] / 100 if item["amount_cents"] is not None else None
        item["review_reason"] = item["observations"] or item["status"]
        if item["status"] != "READY_FOR_SHEET":
            review_rows.append(item)
        elif item["document_type"] == "deposit":
            deposits[_deposit_group_key(item.get("origin_bank"))].append(item)
        else:
            groups[item["destination_account_key"] or "sin_cuenta"].append(item)
    if source_batch:
        output = settings.export_dir / _safe(source_batch)
    elif len(document_ids) == 1:
        output = settings.export_dir / f"documento_{document_ids[0]}"
    elif document_ids:
        output = settings.export_dir / ("documentos_" + "_".join(map(str, document_ids)))
    else:
        output = settings.export_dir
    _prepare_output(output)
    written: list[str] = []
    for key, items in groups.items():
        path = output / f"transferencias_{_safe(key)}.xlsx"
        _write(path, items)
        written.append(str(path))
    row_trace: list[dict] = []
    for key, items in groups.items():
        filename = f"transferencias_{_safe(key)}.xlsx"
        for excel_row, item in enumerate(items, start=2):
            row_trace.append({
                "document_id": item["document_id"],
                "arrival_sequence": item["arrival_sequence"],
                "original_filename": item["original_filename"],
                "source_system": item["source_system"],
                "source_document_id": item["source_document_id"],
                "source_url": item["source_url"],
                "workbook": filename,
                "row": excel_row,
            })
    for key, items in deposits.items():
        path = output / f"depositos_{_safe(key)}.xlsx"
        _write(path, items, deposit=True)
        written.append(str(path))
        for excel_row, item in enumerate(items, start=2):
            row_trace.append({
                "document_id": item["document_id"],
                "arrival_sequence": item["arrival_sequence"],
                "original_filename": item["original_filename"],
                "source_system": item["source_system"],
                "source_document_id": item["source_document_id"],
                "source_url": item["source_url"],
                "workbook": path.name,
                "row": excel_row,
            })
    if review_rows:
        path = output / "REVIEW.xlsx"
        _write(path, review_rows, review=True)
        written.append(str(path))
        for excel_row, item in enumerate(review_rows, start=2):
            row_trace.append({
                "document_id": item["document_id"],
                "arrival_sequence": item["arrival_sequence"],
                "original_filename": item["original_filename"],
                "source_system": item["source_system"],
                "source_document_id": item["source_document_id"],
                "source_url": item["source_url"],
                "workbook": path.name,
                "row": excel_row,
            })
    trace_by_document = {item["document_id"]: item for item in row_trace}
    control_path = output / "CONTROL_LOTE.xlsx"
    _write_control(control_path, [dict(item, amount=item["amount_cents"] / 100 if item["amount_cents"] is not None else None, review_reason=item["observations"] or item["status"]) for item in map(dict, rows)], trace_by_document)
    written.append(str(control_path))
    report = {
        "generated_at": datetime.now(ZoneInfo("America/Argentina/Buenos_Aires")).date().isoformat(),
        "source_batch": source_batch,
        "document_id": document_ids[0] if len(document_ids) == 1 else None,
        "document_ids": document_ids,
        "files": written,
        "rows": len(rows),
        "status_counts": dict(Counter(item["status"] for item in map(dict, rows))),
        "row_trace": row_trace,
    }
    output.mkdir(parents=True, exist_ok=True)
    (output / "manifest.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report
