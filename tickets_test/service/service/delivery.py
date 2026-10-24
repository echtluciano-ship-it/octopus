"""Fail-closed delivery validation; no writes or provider calls."""
from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from service.db import connection
from service.object_store import get_original


def receipt_integrity() -> dict:
    # Check each message, not just aggregate totals: a surplus cannot hide a loss.
    # Global scope is intentionally conservative until batches have a durable
    # message-level boundary (source_batch currently belongs to documents).
    with connection() as conn:
        rows = conn.execute("""
            SELECT m.message_id,m.expected_attachments,
                   COUNT(d.document_id) AS stored,
                   COUNT(d.document_id) FILTER (
                     WHERE d.attachment_index < 0 OR
                           d.attachment_index >= m.expected_attachments) AS invalid_indices
            FROM inbound_messages m
            LEFT JOIN ticket_documents d USING(message_id)
            GROUP BY m.message_id,m.expected_attachments
        """).fetchall()
        failures = conn.execute("""SELECT COUNT(*) AS count FROM ingest_failures
            WHERE resolved_at IS NULL""").fetchone()["count"]
    mismatches = [int(r["message_id"]) for r in rows
                  if r["expected_attachments"] != r["stored"] or r["invalid_indices"]]
    return {"ready": not mismatches and failures == 0,
            "expected": sum(r["expected_attachments"] for r in rows),
            "stored": sum(r["stored"] for r in rows),
            "mismatched_messages": mismatches, "unresolved_ingest_errors": failures}


def _normal(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None if value == "" else value


def validate_delivery(manifest_path: Path, attachments: list[Path]) -> dict:
    from service.export_excel import (
        ERIKA_COLUMNS, _erika_observation, evaluate_delivery_gate, load_delivery_rows,
    )

    report = json.loads(manifest_path.read_text(encoding="utf-8"))
    if report.get("delivery_gate", {}).get("delivery_ready") is not True:
        raise ValueError("Delivery gate is missing or blocked")
    integrity = receipt_integrity()
    if not integrity["ready"]:
        raise ValueError("Expected attachments do not reconcile with stored originals")
    rows = load_delivery_rows(report.get("source_batch"), report.get("document_ids"))
    gate = evaluate_delivery_gate(rows)
    if not gate["delivery_ready"]:
        raise ValueError("Current document states block delivery")
    requested = report.get("document_ids") or []
    if requested and (len(requested) != len(set(requested)) or
                      set(requested) != {r["document_id"] for r in rows}):
        raise ValueError("Requested documents are missing or duplicated")
    groups = defaultdict(list)
    for row in rows:
        groups[str(row["group_id"])].append(row)
    books = report.get("erika_workbooks", [])
    if len(books) != len(groups) or {b["group_id"] for b in books} != set(groups):
        raise ValueError("Exactly one workbook per group is required")
    paths = [Path(p).resolve() for p in attachments]
    names = [Path(b["path"]).name for b in books]
    if (len(paths) != len(books) or len(set(names)) != len(names) or
            len(set(paths)) != len(paths) or {p.name for p in paths} != set(names)):
        raise ValueError("Attachments must be exactly the group workbooks")
    expected_trace = []
    for book in books:
        group_rows = groups[book["group_id"]]
        path = next(p for p in paths if p.name == Path(book["path"]).name)
        if path.suffix.lower() != ".xlsx" or book["rows"] != len(group_rows):
            raise ValueError("Invalid workbook or row count")
        workbook = load_workbook(path, data_only=False)
        try:
            if workbook.sheetnames != ["Tickets"] or workbook.security.lockStructure:
                raise ValueError("Workbook structure must be editable and contain only Tickets")
            sheet = workbook["Tickets"]
            if sheet.protection.sheet or sheet.max_column != 7:
                raise ValueError("Workbook is protected or has unexpected columns")
            values = list(sheet.values)
            if values[0] != ERIKA_COLUMNS or len(values) != len(group_rows) + 1:
                raise ValueError("Workbook columns or row count differ from ledger")
            for index, (actual, row) in enumerate(zip(values[1:], group_rows), start=2):
                expected = (row.get("ticket_date"), row.get("amount"),
                            row.get("identifier_value"), row.get("full_name"),
                            row.get("tax_id"), row.get("origin_bank"), _erika_observation(row))
                if tuple(map(_normal, actual)) != tuple(map(_normal, expected)):
                    raise ValueError("Workbook data or arrival order differ from ledger")
                expected_trace.append({"document_id": row["document_id"],
                    "group_id": row["group_id"], "arrival_sequence": row["arrival_sequence"],
                    "attachment_index": row["attachment_index"],
                    "workbook": path.name, "row": index})
        finally:
            workbook.close()
    if report.get("erika_row_trace") != expected_trace or report.get("rows") != len(rows):
        raise ValueError("Manifest does not match the current ledger")
    with connection() as conn:
        originals = conn.execute("""SELECT document_id,object_key,sha256,size_bytes
            FROM ticket_documents WHERE document_id = ANY(%s)""",
            ([r["document_id"] for r in rows],)).fetchall()
    if len(originals) != len(rows):
        raise ValueError("Missing original records")
    for original in originals:
        content = get_original(original["object_key"])
        if (len(content) != original["size_bytes"] or
                hashlib.sha256(content).hexdigest() != original["sha256"]):
            raise ValueError("Original content failed integrity verification")
    return gate
