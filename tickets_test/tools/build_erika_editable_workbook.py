from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


HEADERS = (
    "FECHA TICKET",
    "MONTO",
    "NRO COMPROBANTE",
    "NOMBRE Y APELLIDO",
    "CUIT",
    "BANCO",
    "OBSERVACIONES",
)


def _observation(row: tuple) -> str:
    status = str(row[2] or "")
    if status == "READY_FOR_SHEET":
        return ""
    if status == "WRONG_ACCOUNT_REVIEW":
        return "REVIEW: la cuenta de destino no corresponde a una cuenta válida."
    if not row[3]:
        return "REVIEW: no se pudo confirmar si corresponde a una transferencia o a un depósito."
    fields = (
        (4, "fecha del ticket"),
        (5, "monto"),
        (7, "número de comprobante, operación o ID Coelsa"),
        (8, "nombre y apellido"),
        (9, "CUIT"),
        (10, "banco"),
    )
    missing = [label for index, label in fields if row[index] in (None, "")]
    if missing:
        return f"REVIEW: falta confirmar {', '.join(missing)}."
    return "REVIEW: el comprobante requiere validación manual."


def build(source: Path, output: Path) -> None:
    source_book = load_workbook(source, data_only=True, read_only=True)
    source_sheet = source_book["Control"]
    source_rows = list(source_sheet.iter_rows(min_row=2, max_row=21, values_only=True))
    if len(source_rows) != 20:
        raise ValueError(f"Se esperaban 20 comprobantes y se encontraron {len(source_rows)}")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tickets"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(HEADERS)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    ready_fill = PatternFill("solid", fgColor="E6F4EA")
    review_fill = PatternFill("solid", fgColor="FCE8E6")
    review_font = Font(name="Arial", size=10, color="A61B1B")
    body_font = Font(name="Arial", size=10, color="1F2937")
    border = Border(bottom=Side(style="thin", color="5BC0DE"))

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for source_row in source_rows:
        status = str(source_row[2] or "")
        sheet.append(
            (
                source_row[4],
                source_row[5],
                None if source_row[7] in (None, "") else str(source_row[7]),
                source_row[8],
                source_row[9],
                source_row[10],
                _observation(source_row),
            )
        )
        row_number = sheet.max_row
        review = status != "READY_FOR_SHEET"
        for cell in sheet[row_number]:
            cell.fill = review_fill if review else ready_fill
            cell.font = review_font if review else body_font
            cell.border = border
            cell.alignment = Alignment(vertical="center")
        sheet.cell(row_number, 1).alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(row_number, 2).alignment = Alignment(horizontal="right", vertical="center")
        sheet.cell(row_number, 3).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.cell(row_number, 3).number_format = "@"
        sheet.cell(row_number, 5).alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(row_number, 6).alignment = Alignment(horizontal="center", vertical="center")
        sheet.cell(row_number, 7).alignment = Alignment(vertical="center", wrap_text=True)

    sheet.auto_filter.ref = "A1:G21"
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 17
    sheet.column_dimensions["C"].width = 40
    sheet.column_dimensions["D"].width = 31
    sheet.column_dimensions["E"].width = 19
    sheet.column_dimensions["F"].width = 20
    sheet.column_dimensions["G"].width = 58
    sheet.row_dimensions[1].height = 25
    for row_number in range(2, 22):
        sheet.row_dimensions[row_number].height = 34
        sheet.cell(row_number, 1).number_format = "dd/mm/yyyy"
        sheet.cell(row_number, 2).number_format = '"$"#,##0.00'

    sheet.protection.sheet = False
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Uso: python build_erika_editable_workbook.py <control.xlsx> <salida.xlsx>")
    build(Path(sys.argv[1]), Path(sys.argv[2]))
