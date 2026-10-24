from __future__ import annotations

import hashlib
import json
import sqlite3
from collections import Counter
from contextlib import closing
from datetime import date
from io import BytesIO
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook

from data_loader import clean_text, parse_date, parse_decimal, parse_period


def sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def download_google_sheet_xlsx(drive_id: str, timeout: int = 120) -> bytes:
    url = f"https://docs.google.com/spreadsheets/d/{drive_id}/export?format=xlsx"
    request = Request(url, headers={"User-Agent": "OCTOPUS billing freshness check"})
    with urlopen(request, timeout=timeout) as response:
        return response.read()


def verify_live_billing_source(app_dir: Path, fetch=download_google_sheet_xlsx) -> str:
    metadata_path = app_dir / "data" / "billing_source.json"
    source_path = app_dir / "data" / "FACTURACION OCTOPUS.xlsx"
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    live_hash = sha256_bytes(fetch(metadata["drive_id"]))
    local_hash = sha256_bytes(source_path.read_bytes())
    recorded_hash = metadata["sha256"]
    if live_hash != local_hash or local_hash != recorded_hash:
        raise ValueError(
            "Facturacion Historica cambio durante la actualizacion: "
            f"live={live_hash}, local={local_hash}, recorded={recorded_hash}. "
            "Descargar nuevamente, reconstruir la base y repetir los controles antes de publicar."
        )
    return live_hash


def _source_rows_for_load_date(source_path: Path, load_date: date) -> Counter:
    workbook = load_workbook(BytesIO(source_path.read_bytes()), data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    expected = Counter()
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        values = list(raw) + [None] * 13
        if parse_date(values[0]) != load_date:
            continue
        client = clean_text(values[2])
        period = parse_period(values[5])
        net = parse_decimal(values[6])
        total = parse_decimal(values[10])
        if not client or period is None or (net is None and total is None):
            continue
        expected[(client, clean_text(values[4]), net, total)] += 1
    return expected


def _database_rows_for_load_date(conn: sqlite3.Connection, load_date: date) -> Counter:
    rows = conn.execute(
        """SELECT original_client_name, invoice_number, net_amount, total_amount
        FROM billing_operations WHERE load_date = ?""",
        (load_date.isoformat(),),
    )
    return Counter((clean_text(client), clean_text(invoice), net, total) for client, invoice, net, total in rows)


def reconcile_billing_load_date(app_dir: Path, load_date: date) -> dict:
    source_rows = _source_rows_for_load_date(app_dir / "data" / "FACTURACION OCTOPUS.xlsx", load_date)
    with closing(sqlite3.connect(app_dir / "octopus.db")) as conn:
        database_rows = _database_rows_for_load_date(conn, load_date)
    if source_rows != database_rows:
        missing = list((source_rows - database_rows).elements())
        unexpected = list((database_rows - source_rows).elements())
        raise ValueError(
            f"Facturacion del {load_date.isoformat()} no reconcilia: "
            f"faltantes={missing}, inesperadas={unexpected}"
        )
    return {
        "load_date": load_date.isoformat(),
        "rows": sum(source_rows.values()),
        "net_amount": round(sum((row[2] or 0) * count for row, count in source_rows.items()), 2),
    }
