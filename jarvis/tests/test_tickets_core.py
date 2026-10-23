from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from jarvis.tickets.excel import write_editable_sheet
from jarvis.tickets.models import AttachmentInput, MessageInput, TicketStatus
from jarvis.tickets.rules import TRANSFER_COLUMNS, close_window
from jarvis.tickets.store import TicketInbox


class TicketsCoreTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.inbox = TicketInbox(self.root / "tickets_test.db")

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _message(self, number: int, group: str = "Tickets A") -> MessageInput:
        return MessageInput(
            provider="fixture",
            instance_id="octopus-test",
            group_id=group,
            group_name=group,
            provider_message_id=f"msg-{number}",
            sender_id=f"sender-{number}",
            sender_name="Persona Test",
            message_timestamp=f"2026-10-20T14:{number:02d}:00-03:00",
            received_at=f"2026-10-20T14:{number:02d}:01-03:00",
            raw_payload_json=json.dumps({"fixture": number}),
            expected_attachments=1,
        )

    def _attachment(self, content: bytes, index: int = 0) -> AttachmentInput:
        digest = hashlib.sha256(content).hexdigest()
        return AttachmentInput(
            attachment_index=index,
            original_filename=f"{digest[:8]}.jpg",
            mime_type="image/jpeg",
            size_bytes=len(content),
            sha256=digest,
            object_uri=f"fixture://{digest}",
        )

    def test_replayed_webhook_is_idempotent(self) -> None:
        first = self.inbox.record_message(self._message(1), [self._attachment(b"ticket-one")])
        second = self.inbox.record_message(self._message(1), [self._attachment(b"ticket-one")])
        self.assertTrue(first.inserted)
        self.assertFalse(second.inserted)
        self.assertEqual(first.document_ids, second.document_ids)
        self.assertEqual(len(self.inbox.documents_in_arrival_order()), 1)

    def test_event_is_visible_while_media_download_is_missing_and_can_recover(self) -> None:
        message = self._message(1)
        first = self.inbox.record_message(message, [])
        missing = self.inbox.reconciliation()
        self.assertTrue(first.inserted)
        self.assertEqual(missing["received_messages"], 1)
        self.assertEqual(missing["missing_documents"], 1)
        self.assertFalse(missing["ingest_balanced"])

        retry = self.inbox.record_message(message, [self._attachment(b"downloaded-later")])
        recovered = self.inbox.reconciliation()
        self.assertFalse(retry.inserted)
        self.assertEqual(len(retry.document_ids), 1)
        self.assertEqual(recovered["missing_documents"], 0)
        self.assertTrue(recovered["ingest_balanced"])

    def test_exact_file_across_groups_is_confirmed_duplicate(self) -> None:
        original = self.inbox.record_message(self._message(1), [self._attachment(b"same-binary")])
        duplicate = self.inbox.record_message(
            self._message(2, "Tickets B"), [self._attachment(b"same-binary")]
        )
        self.assertEqual(duplicate.exact_duplicate_ids, duplicate.document_ids)
        rows = self.inbox.documents_in_arrival_order()
        self.assertEqual(rows[1]["duplicate_of"], original.document_ids[0])
        self.assertEqual(rows[1]["duplicate_reason"], "identical SHA-256 and byte size")

    def test_same_person_amount_and_date_are_not_duplicate_evidence(self) -> None:
        first = self.inbox.record_message(self._message(1), [self._attachment(b"operation-a")])
        second = self.inbox.record_message(self._message(2), [self._attachment(b"operation-b")])
        self.assertFalse(first.exact_duplicate_ids)
        self.assertFalse(second.exact_duplicate_ids)
        self.assertEqual(len(self.inbox.documents_in_arrival_order()), 2)

    def test_arrival_order_is_not_business_date_order(self) -> None:
        first = self.inbox.record_message(self._message(1), [self._attachment(b"later-date")])
        second = self.inbox.record_message(self._message(2), [self._attachment(b"earlier-date")])
        rows = self.inbox.documents_in_arrival_order()
        self.assertEqual([row["document_id"] for row in rows], [first.document_ids[0], second.document_ids[0]])

    def test_reconciliation_never_loses_a_document(self) -> None:
        results = []
        for number in range(1, 64):
            results.append(
                self.inbox.record_message(self._message(number), [self._attachment(f"ticket-{number}".encode())])
            )
        for result in results[:57]:
            self.inbox.transition(
                result.document_ids[0], TicketStatus.READY_FOR_SHEET,
                actor="test", reason="valid extraction"
            )
        for result in results[57:60]:
            self.inbox.transition(
                result.document_ids[0], TicketStatus.DUPLICATE_CONFIRMED,
                actor="human-test", reason="confirmed duplicate"
            )
        for result in results[60:62]:
            self.inbox.transition(
                result.document_ids[0], TicketStatus.REVIEW,
                actor="test", reason="insufficient evidence"
            )
        self.inbox.transition(
            results[62].document_ids[0], TicketStatus.WRONG_ACCOUNT_REVIEW,
            actor="test", reason="destination account is not configured"
        )
        report = self.inbox.reconciliation()
        self.assertEqual(report["received_documents"], 63)
        self.assertEqual(report["processed"], 57)
        self.assertEqual(report["duplicates"], 3)
        self.assertEqual(report["review"], 3)
        self.assertEqual(report["in_progress"], 0)
        self.assertTrue(report["balanced"])

    def test_close_windows_are_half_open_and_cover_weekend(self) -> None:
        friday = close_window(date(2026, 10, 23))
        monday = close_window(date(2026, 10, 26))
        tuesday = close_window(date(2026, 10, 27))
        self.assertEqual(friday[0].isoformat(), "2026-10-22T15:00:00-03:00")
        self.assertEqual(monday[0].isoformat(), "2026-10-23T15:00:00-03:00")
        self.assertEqual(monday[1], tuesday[0])

    def test_holiday_extends_window_to_next_business_day(self) -> None:
        tuesday = close_window(
            date(2026, 10, 13),
            holidays={date(2026, 10, 12)},
        )
        self.assertEqual(tuesday[0].isoformat(), "2026-10-09T15:00:00-03:00")
        self.assertEqual(tuesday[1].isoformat(), "2026-10-13T15:00:00-03:00")

    def test_close_is_not_allowed_on_holiday(self) -> None:
        with self.assertRaises(ValueError):
            close_window(date(2026, 10, 12), holidays={date(2026, 10, 12)})

    def test_generated_workbook_matches_reference_columns_and_is_editable(self) -> None:
        output = self.root / "tickets.xlsx"
        write_editable_sheet(
            output,
            [{
                "FECHA TICHET": date(2026, 10, 21),
                "MONTO": 123456,
                "NRO COMPROBANTE": "OP-1",
                "NOMBRE Y APELLIDO": "Persona Test",
                "CUIT": "20-00000000-0",
                "BANCO": "Banco Test",
            }],
            document_type="transfer",
        )
        workbook = load_workbook(output)
        sheet = workbook["Hoja1"]
        self.assertEqual(tuple(cell.value for cell in sheet[1]), TRANSFER_COLUMNS)
        sheet["D2"] = "Corregido por Erika"
        workbook.save(output)
        self.assertEqual(load_workbook(output)["Hoja1"]["D2"].value, "Corregido por Erika")


if __name__ == "__main__":
    unittest.main()
